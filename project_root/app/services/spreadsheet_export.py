"""Spreadsheet export service for generating Excel project schedules."""
import os
import uuid
import datetime
from flask import current_app

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def create_project_schedule(project_data, project_name):
    """Create an Excel spreadsheet schedule from project data."""
    if OPENPYXL_AVAILABLE:
        return generate_excel_file(project_data, project_name)
    else:
        return generate_csv_file(project_data, project_name)


def generate_excel_file(project_data, project_name):
    """Generate a formatted Excel file with project schedule."""
    wb = Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # ===== TASKS SHEET =====
    ws_tasks = wb.active
    ws_tasks.title = "Tasks"

    # Task headers
    task_headers = ["ID", "Task Name", "Description", "Duration (days)", "Start Day", "End Day", "Predecessors", "Assigned Resources"]
    for col, header in enumerate(task_headers, 1):
        cell = ws_tasks.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Calculate task schedule (early start/finish)
    task_schedule = calculate_schedule(project_data["tasks"])

    # Task data
    for row, task in enumerate(project_data["tasks"], 2):
        task_id = task["id"]
        schedule = task_schedule.get(task_id, {"start": 0, "finish": 0})

        ws_tasks.cell(row=row, column=1, value=task_id).border = thin_border
        ws_tasks.cell(row=row, column=2, value=task["name"]).border = thin_border
        ws_tasks.cell(row=row, column=3, value=task.get("description", "")).border = thin_border
        ws_tasks.cell(row=row, column=4, value=task["duration"]).border = thin_border
        ws_tasks.cell(row=row, column=5, value=schedule["start"]).border = thin_border
        ws_tasks.cell(row=row, column=6, value=schedule["finish"]).border = thin_border

        # Predecessors
        predecessors = ", ".join(str(p) for p in task.get("predecessors", []))
        ws_tasks.cell(row=row, column=7, value=predecessors if predecessors else "-").border = thin_border

        # Resources
        resources = ", ".join(task.get("resources", []))
        ws_tasks.cell(row=row, column=8, value=resources if resources else "-").border = thin_border

    # Adjust column widths for tasks sheet
    ws_tasks.column_dimensions['A'].width = 8
    ws_tasks.column_dimensions['B'].width = 35
    ws_tasks.column_dimensions['C'].width = 40
    ws_tasks.column_dimensions['D'].width = 15
    ws_tasks.column_dimensions['E'].width = 12
    ws_tasks.column_dimensions['F'].width = 12
    ws_tasks.column_dimensions['G'].width = 15
    ws_tasks.column_dimensions['H'].width = 30

    # ===== RESOURCES SHEET =====
    ws_resources = wb.create_sheet(title="Resources")

    # Resource headers
    resource_headers = ["ID", "Resource Name", "Capacity (%)", "Assigned Tasks"]
    for col, header in enumerate(resource_headers, 1):
        cell = ws_resources.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Build resource-to-tasks mapping
    resource_tasks = {}
    for task in project_data["tasks"]:
        for res_name in task.get("resources", []):
            if res_name not in resource_tasks:
                resource_tasks[res_name] = []
            resource_tasks[res_name].append(task["name"])

    # Resource data
    for row, resource in enumerate(project_data["resources"], 2):
        ws_resources.cell(row=row, column=1, value=resource["id"]).border = thin_border
        ws_resources.cell(row=row, column=2, value=resource["name"]).border = thin_border
        ws_resources.cell(row=row, column=3, value=resource.get("capacity", 100)).border = thin_border

        # Assigned tasks
        assigned = resource_tasks.get(resource["name"], [])
        ws_resources.cell(row=row, column=4, value=", ".join(assigned) if assigned else "-").border = thin_border

    # Adjust column widths for resources sheet
    ws_resources.column_dimensions['A'].width = 8
    ws_resources.column_dimensions['B'].width = 25
    ws_resources.column_dimensions['C'].width = 15
    ws_resources.column_dimensions['D'].width = 50

    # ===== TIMELINE SHEET (Gantt-style) =====
    ws_timeline = wb.create_sheet(title="Timeline")

    # Calculate project duration
    max_duration = max((s["finish"] for s in task_schedule.values()), default=0)

    # Timeline headers
    ws_timeline.cell(row=1, column=1, value="Task Name")
    ws_timeline.cell(row=1, column=1).font = header_font
    ws_timeline.cell(row=1, column=1).fill = header_fill
    ws_timeline.cell(row=1, column=1).border = thin_border

    # Day columns (limit to 60 days for readability)
    display_days = min(max_duration + 5, 60)
    for day in range(1, display_days + 1):
        cell = ws_timeline.cell(row=1, column=day + 1, value=f"Day {day}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        ws_timeline.column_dimensions[get_column_letter(day + 1)].width = 6

    ws_timeline.column_dimensions['A'].width = 30

    # Task bars
    task_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")

    for row, task in enumerate(project_data["tasks"], 2):
        task_id = task["id"]
        schedule = task_schedule.get(task_id, {"start": 0, "finish": 0})

        ws_timeline.cell(row=row, column=1, value=task["name"]).border = thin_border

        # Fill cells for task duration
        start_day = schedule["start"] + 1  # Convert to 1-indexed
        end_day = schedule["finish"]

        for day in range(1, display_days + 1):
            cell = ws_timeline.cell(row=row, column=day + 1)
            cell.border = thin_border
            if start_day <= day <= end_day:
                cell.fill = task_fill

    # ===== SUMMARY SHEET =====
    ws_summary = wb.create_sheet(title="Summary")

    summary_data = [
        ("Project Name", project_name),
        ("Created", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("Total Tasks", len(project_data["tasks"])),
        ("Total Resources", len(project_data["resources"])),
        ("Project Duration (days)", max_duration),
    ]

    for row, (label, value) in enumerate(summary_data, 1):
        label_cell = ws_summary.cell(row=row, column=1, value=label)
        label_cell.font = Font(bold=True)
        ws_summary.cell(row=row, column=2, value=value)

    ws_summary.column_dimensions['A'].width = 25
    ws_summary.column_dimensions['B'].width = 30

    # Create output directory if it doesn't exist
    output_dir = os.path.join(current_app.instance_path, 'generated')
    os.makedirs(output_dir, exist_ok=True)

    # Create the file path
    file_name = f"{project_name.replace(' ', '_')}_{uuid.uuid4()}.xlsx"
    file_path = os.path.join(output_dir, file_name)

    # Save workbook
    wb.save(file_path)

    return file_path


def generate_csv_file(project_data, project_name):
    """Fallback: Generate a CSV file if openpyxl is not available."""
    import csv

    # Create output directory if it doesn't exist
    output_dir = os.path.join(current_app.instance_path, 'generated')
    os.makedirs(output_dir, exist_ok=True)

    # Create the file path
    file_name = f"{project_name.replace(' ', '_')}_{uuid.uuid4()}.csv"
    file_path = os.path.join(output_dir, file_name)

    # Calculate schedule
    task_schedule = calculate_schedule(project_data["tasks"])

    with open(file_path, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)

        # Header
        writer.writerow(["ID", "Task Name", "Description", "Duration (days)", "Start Day", "End Day", "Predecessors", "Assigned Resources"])

        # Tasks
        for task in project_data["tasks"]:
            task_id = task["id"]
            schedule = task_schedule.get(task_id, {"start": 0, "finish": 0})

            predecessors = ", ".join(str(p) for p in task.get("predecessors", []))
            resources = ", ".join(task.get("resources", []))

            writer.writerow([
                task_id,
                task["name"],
                task.get("description", ""),
                task["duration"],
                schedule["start"],
                schedule["finish"],
                predecessors if predecessors else "-",
                resources if resources else "-"
            ])

        # Blank row
        writer.writerow([])

        # Resources section
        writer.writerow(["Resources"])
        writer.writerow(["ID", "Resource Name", "Capacity (%)"])
        for resource in project_data["resources"]:
            writer.writerow([
                resource["id"],
                resource["name"],
                resource.get("capacity", 100)
            ])

    return file_path


def calculate_schedule(tasks):
    """Calculate early start and early finish times for all tasks."""
    task_dict = {task["id"]: task for task in tasks}
    schedule = {}

    # Initialize all tasks
    for task in tasks:
        schedule[task["id"]] = {"start": 0, "finish": 0}

    # Forward pass - calculate early start and finish
    for task in tasks:
        task_id = task["id"]
        predecessors = task.get("predecessors", [])

        if not predecessors:
            # No predecessors, start at day 0
            schedule[task_id]["start"] = 0
        else:
            # Start after all predecessors finish
            max_pred_finish = 0
            for pred_id in predecessors:
                if pred_id in schedule:
                    pred_finish = schedule[pred_id]["finish"]
                    if pred_finish > max_pred_finish:
                        max_pred_finish = pred_finish
            schedule[task_id]["start"] = max_pred_finish

        schedule[task_id]["finish"] = schedule[task_id]["start"] + task["duration"]

    return schedule
